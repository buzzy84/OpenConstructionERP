"""Synthetic, deterministic tests for the project Drawing Log (never use customer data)."""
from datetime import date
from io import BytesIO
from uuid import uuid4

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from app.modules.production_register.excel_io import COLUMN_HEADERS, _boq_refs, _write_text, parse_xlsx
from app.modules.production_register.manifest import manifest
from app.modules.production_register.models import (
    ProductionRegister, ProductionItem, ItemBOQLink, ItemRelation, WorkstreamDefinition,
    ItemCheckpoint, ItemReference, RegisterAudit,
)
from app.modules.production_register.schemas import ItemInput, ItemPatch, CheckpointWrite, ReferenceCreate
from app.modules.production_register.service import INITIAL_WORKSTREAMS, RegisterService


def sample_workbook(*, formulas=False):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Foglio1'
    sheet['A1'] = 'Synthetic drawing log'
    for index, h in enumerate(COLUMN_HEADERS, start=1):
        sheet.cell(3, index, h)
    sheet['A4'] = 'BASEMENT FLOOR'
    sheet['A5'] = 'WALL FINISHES'
    sheet['A6'] = 'ROOM-01'
    sheet['B6'] = 'A-001\nB-002;C-003'
    sheet['C6'] = 'UNIT-1+UNIT-2'
    sheet['D6'] = 'Display wall'
    sheet['F6'] = 'a corpo'
    sheet['G6'] = '4ml'
    sheet['K6'] = date(2026,8,25)
    sheet['L6'] = 'Supplier pending'
    sheet['A7'] = 'ROOM-02'
    sheet['B7'] = 'A-001'
    sheet['C7'] = 'UNIT-1+UNIT-2'
    sheet['D7'] = 'Second distinct occurrence'
    sheet['A8'] = 'ROOM-02'
    sheet['C8'] = 'UNIT-1+UNIT-2'
    sheet['D8'] = 'Third distinct occurrence'
    if formulas:
        sheet['D8'] = '=1+1'
    data=BytesIO();book.save(data)
    return data.getvalue()


def test_manifest_and_schema_expose_distinct_eight_tables():
    assert manifest.name == 'oe_production_register'
    assert manifest.category == 'community'
    classes = (ProductionRegister,ProductionItem,ItemBOQLink,ItemRelation,WorkstreamDefinition,
               ItemCheckpoint,ItemReference,RegisterAudit)
    assert len({x.__tablename__ for x in classes}) == 8
    assert all(x.__tablename__.startswith('oe_production_register_') for x in classes)
    assert 'item_code' not in [col.name for c in ProductionItem.__table__.constraints if hasattr(c,'columns') and c.__class__.__name__=='UniqueConstraint' for col in c.columns]
    assert len(INITIAL_WORKSTREAMS) == 8
    assert {x[0] for x in INITIAL_WORKSTREAMS} >= {'shipment_1','shipment_2','production','glass'}


def test_excel_keeps_many_boq_references_and_duplicate_item_codes():
    parsed = parse_xlsx(sample_workbook(), 'synthetic.xlsx')
    assert len(parsed['rows']) == 3
    a,b,c = parsed['rows']
    assert a['floor'] == 'BASEMENT FLOOR'
    assert a['category'] == 'WALL FINISHES'
    assert a['boq_refs'] == ['A-001','B-002','C-003']
    assert b['boq_refs'] == ['A-001']
    assert a['item_code'] == b['item_code'] == c['item_code'] == 'UNIT-1+UNIT-2'
    assert a['quoted_qty_raw'] == 'a corpo'
    assert a['production_qty_raw'] == '4ml'
    assert a['checkpoint_raw']['production'] == 'Supplier pending'
    assert a['reported_approval_date'] == '2026-08-25'


def test_excel_formula_refusal_and_magic_bytes():
    with pytest.raises(HTTPException) as exc:
        parse_xlsx(sample_workbook(formulas=True), 'example.xlsx')
    assert exc.value.status_code == 422
    with pytest.raises(HTTPException):
        parse_xlsx(b'not a zip', 'example.xlsx')
    with pytest.raises(HTTPException):
        parse_xlsx(sample_workbook(), 'example.xlsm')


def test_xlsx_formula_like_strings_are_text_not_formulas():
    book=Workbook(); ws=book.active
    tokens=['=1+1','+1+2','-1','@LOOKUP','UNIT-1+UNIT-2']
    for i,text in enumerate(tokens,1):
        _write_text(ws.cell(i,1),text)
    stream=BytesIO();book.save(stream);stream.seek(0)
    loaded=load_workbook(stream)
    for i,expected in enumerate(tokens,1):
        cell=loaded.active.cell(i,1)
        assert cell.value == expected
        assert cell.data_type != 'f'


def test_schema_does_not_accept_cross_project_or_audit_spoofing():
    with pytest.raises(ValidationError):
        ItemInput(item_code='X-1',project_id=str(uuid4()))
    with pytest.raises(ValidationError):
        ItemInput(item_code='X-1',source='api',created_by=str(uuid4()))
    with pytest.raises(ValidationError):
        ItemPatch(expected_version=0)
    with pytest.raises(ValidationError):
        CheckpointWrite(expected_version=0,status='approved')
    with pytest.raises(ValidationError):
        ReferenceCreate(kind='external',namespace='source',external_key=None)


def test_boq_ref_split_only_on_newline_semicolon():
    assert _boq_refs('CLN1+C14\nGRP_A-02; CLD-03') == ['CLN1+C14','GRP_A-02','CLD-03']


def test_optimistic_check_rejects_stale_clients():
    with pytest.raises(HTTPException) as exc:
        RegisterService.check_version(8,7)
    assert exc.value.status_code == 409
    RegisterService.check_version(8,8)


def test_rejects_unrecognized_sync_sheet_schema_version():
    from app.modules.production_register.excel_io import SYNC_SCHEMA
    book = load_workbook(BytesIO(sample_workbook()))
    sh = book.create_sheet('_OCERP_SYNC')
    sh.append(['row_number', 'project_id', 'register_id', 'item_id', 'item_version', 'row_hash', 'checkpoint_data', 'visible_hashes', 'original_identity', 'schema_version'])
    sh.append([6, str(uuid4()), str(uuid4()), str(uuid4()), 1, '', '{}', '{}', '{}', SYNC_SCHEMA + 999])
    out = BytesIO(); book.save(out)
    with pytest.raises(HTTPException) as exc:
        parse_xlsx(out.getvalue(), 'synthetic.xlsx')
    assert exc.value.status_code == 422
    assert exc.value.detail['code'] == 'INVALID_SYNC_SCHEMA'


def test_sync_from_another_register_cannot_be_overridden():
    from app.modules.production_register.excel_io import validate_resolution_override
    with pytest.raises(HTTPException) as exc:
        validate_resolution_override('WRONG_REGISTER', {'action': 'create'})
    assert exc.value.status_code == 409
    with pytest.raises(HTTPException):
        validate_resolution_override('ITEM_VERSION_CONFLICT', {'action': 'match', 'item_id': str(uuid4())})
    validate_resolution_override('AMBIGUOUS_LEGACY_MATCH', {'action': 'create'})


def test_safeguards_archived_records_against_writes():
    from app.modules.production_register.service import assert_item_editable
    class Item:
        lifecycle = 'archived'
    with pytest.raises(HTTPException) as exc:
        assert_item_editable(Item())
    assert exc.value.status_code == 409
    Item.lifecycle = 'active'
    assert_item_editable(Item())


def test_reimport_blank_clears_previous_raw_checkpoint_value():
    from types import SimpleNamespace
    from app.modules.production_register.excel_io import needs_checkpoint_update
    assert needs_checkpoint_update(SimpleNamespace(reported_value='Supplier pending'), None)
    assert not needs_checkpoint_update(None, None)
    assert needs_checkpoint_update(None, 'new raw value')
    assert not needs_checkpoint_update(SimpleNamespace(reported_value='unchanged'), 'unchanged')


def test_structured_checkpoint_survives_unchanged_xlsx_cell():
    from app.modules.production_register.excel_io import _hash
    book = load_workbook(BytesIO(sample_workbook()))
    ws = book.active
    ws['L6'] = 'Supplier pending'
    sh = book.create_sheet('_OCERP_SYNC')
    sh.append(['row_number','project_id','register_id','item_id','item_version','row_hash','checkpoint_data','visible_hashes','original_identity','schema_version'])
    stored = {
       'production': {'status':'on_hold','reported_value':'Supplier pending','notes':'Waiting for stone',
                      'planned_date':'2026-09-10','actual_date':None,'assignee_name':'External vendor','version':2}
    }
    import json
    sh.append([6,str(uuid4()),str(uuid4()),str(uuid4()),2,'',json.dumps(stored),
               json.dumps({'production':_hash('Supplier pending')}),'{}',1])
    out=BytesIO();book.save(out)
    result=parse_xlsx(out.getvalue(),'synthetic.xlsx')
    assert result['rows'][0]['checkpoint_structured']['production']['status']=='on_hold'
    assert result['rows'][0]['checkpoint_structured']['production']['notes']=='Waiting for stone'


def test_legacy_export_preserves_structured_checkpoint_on_reimport():
    """An exported status must not be treated as free-text progress on reimport."""
    import asyncio
    from types import SimpleNamespace
    from app.modules.production_register.excel_io import export_xlsx
    reg_id, project_id, item_id, ws_id = (uuid4() for _ in range(4))
    row = ProductionItem(id=item_id, register_id=reg_id, project_id=project_id, display_order=1,
                         floor='BASEMENT FLOOR', room='1', category='WALL', item_code='CLD1',
                         description='Wall', specification='', quoted_qty_raw='a corpo',
                         production_qty_raw='4ml', comments='', change_notes='', notes='',
                         reported_approval_date=None, version=2)
    stream = WorkstreamDefinition(id=ws_id, project_id=project_id, register_id=reg_id,
                                  code='production', label='Production', enabled=True)
    checkpoint = ItemCheckpoint(id=uuid4(), item_id=item_id, workstream_id=ws_id,
                                project_id=project_id, status='completed', reported_value=None,
                                notes='Finished in production', source='manual', version=3)
    reg = ProductionRegister(id=reg_id, project_id=project_id, legacy_code_header='BUZZONI CODE')
    class Repo:
        async def workstreams(self, register_id):return [stream]
        async def links(self, item):return []
        async def checkpoints(self, item):return [checkpoint]
    class FakeDb:
        async def execute(self, statement):
            rows = [row] if statement.column_descriptions[0]['entity'] is ProductionItem else []
            return SimpleNamespace(scalars=lambda:rows)
    svc=SimpleNamespace(db=FakeDb(),repo=Repo(),project_id=project_id)
    blob=asyncio.run(export_xlsx(svc,reg,'legacy'))
    parsed=parse_xlsx(blob,'drawing-log.xlsx')
    line=parsed['rows'][0]
    assert line['item_code']=='CLD1'
    assert line['checkpoint_raw']['production'] is None
    assert line['checkpoint_structured']['production']['status']=='completed'
    assert line['checkpoint_structured']['production']['notes']=='Finished in production'
    assert parsed['legacy_code_header']=='BUZZONI CODE'


def test_import_preview_warns_for_ambiguous_and_unresolved_boq_references():
    from app.modules.production_register.excel_io import reference_warnings
    baseline = [{'ordinal':'A-001','id':str(uuid4())}, {'ordinal':'A-001','id':str(uuid4())},
                {'ordinal':'C-100','id':str(uuid4())}]
    rows = [{'excel_row':6,'boq_refs':['A-001','UNKNOWN','C-100']}]
    warnings = reference_warnings(rows,baseline)
    assert [(x['code'],x['reference']) for x in warnings] == [
      ('AMBIGUOUS_BOQ_REF','A-001'),('UNRESOLVED_BOQ_REF','UNKNOWN')]


def test_readiness_requires_nonblank_signature_and_boq_approval_stamps():
    from app.modules.production_register.service import has_business_stamp
    assert not has_business_stamp(None)
    assert not has_business_stamp('   ')
    assert has_business_stamp('2026-10-27T14:00:00Z')
    assert has_business_stamp(str(uuid4()))


def test_production_register_migration_extends_october_2026_head():
    from pathlib import Path
    revision = (Path(__file__).parents[2] / 'alembic' / 'versions' / 'v3325_production_register.py').read_text()
    assert 'down_revision: Union[str, Sequence[str], None] = "v199_subcontractor_contact_link"' in revision
