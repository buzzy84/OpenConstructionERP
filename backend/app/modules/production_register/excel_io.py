"""Stateless XLSX roundtrip for a twenty-column operational drawing log."""
from __future__ import annotations
import hashlib
import json
import re
from zipfile import ZipFile, BadZipFile
from openpyxl.utils.datetime import from_excel
from collections import Counter, defaultdict
from datetime import date, datetime, time
from io import BytesIO
from uuid import UUID
from typing import Any
from fastapi import HTTPException
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from .models import ProductionItem, ItemBOQLink, ItemCheckpoint
from .schemas import ItemInput, BOQLinkInput, CheckpointWrite
from .service import RegisterService, serialize, fail

COLUMN_HEADERS = (
    "ROOM", "BOQ REF.", "INTERNAL ITEM CODE", "DESCRIZIONE", "SPECIFICHE DI PROGETTO",
    "QUANTITÀ OFFERTA", "QUANTITÀ BUZZONI", "COMMENTI", "CHANGING LOG", "DISEGNATORE",
    "APPROVAZIONE", "PRODUZIONE", "VETRERIA", "ILLUMINAZIONE", "TESSUTI, CARTE, PIETRA",
    "MARMO", "FERRAMENTA", "NOTE", "SPEDIZIONE 1", "SPEDIZIONE 2",
)
LEGACY_WORKSTREAMS = {11: "production", 12: "glass", 13: "lighting", 14: "finishes", 15: "marble",
                      16: "hardware", 18: "shipment_1", 19: "shipment_2"}
MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10000
SYNC_SCHEMA = 1


def validate_resolution_override(issue: str | None, override: dict) -> None:
    """Never turn a foreign-register or stale-version reference into a local write."""
    if issue in ("WRONG_REGISTER", "ITEM_VERSION_CONFLICT"):
        fail(409, issue, "Foreign-register or stale-version sync cannot be overridden")


def _cell(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value)


def _hash(value: Any) -> str:
    return hashlib.sha256(json.dumps(value, sort_keys=True, default=str, ensure_ascii=False).encode()).hexdigest()


def _norm(text: Any) -> str:
    import unicodedata
    unaccent = "".join(c for c in unicodedata.normalize("NFKD", str(text or "")) if not unicodedata.combining(c))
    return re.sub(r"[^A-Z0-9]+", " ", unaccent.upper()).strip()


def _boq_refs(raw: Any) -> list[str]:
    return list(dict.fromkeys(x.strip() for x in re.split(r"[\n;]+", str(raw or "")) if x.strip()))


def _write_text(cell, value: Any) -> None:
    # All string values are explicitly TEXT so Excel never evaluates formula-like input.
    if value is None:
        cell.value = None
    else:
        cell.value, cell.data_type = str(value), "s"


def parse_xlsx(content: bytes, filename: str) -> dict:
    if not filename.lower().endswith(".xlsx") or len(content) > MAX_BYTES or not content.startswith(b"PK"):
        fail(422, "INVALID_XLSX", "Only .xlsx files <= 10 MB are accepted")
    try:
        with ZipFile(BytesIO(content)) as archive:
            total = sum(info.file_size for info in archive.infolist())
            if len(archive.infolist()) > 500 or total > 50 * 1024 * 1024 or any(
                info.file_size > max(info.compress_size, 1) * 200 for info in archive.infolist()
            ):
                fail(422, "XLSX_TOO_LARGE", "XLSX expanded size or compression ratio is excessive")
        book = load_workbook(BytesIO(content), read_only=False, data_only=False)
    except HTTPException:
        raise
    except Exception as exc:
        fail(422, "INVALID_XLSX", f"Unable to read XLSX: {type(exc).__name__}")
    sheets = [s for s in book.worksheets if s.title != "_OCERP_SYNC"]
    if not sheets:
        fail(422, "NO_DRAWING_LOG", "No data worksheet")
    ws = next((s for s in sheets if s.title == "Foglio1"), sheets[0])
    header_row = None
    for idx in range(1, min(ws.max_row, 25) + 1):
        values = [_norm(c.value) for c in ws[idx]]
        if ("ROOM" in values and "BOQ REF" in values and "DESCRIZIONE" in values
                and any(v.endswith("CODE") for v in values)):
            header_row = idx
            break
    if header_row is None:
        fail(422, "HEADERS_NOT_FOUND", "Expected ROOM, BOQ REF., internal CODE, DESCRIZIONE headers")
    mapping = {}
    header_values = [_norm(cell.value) for cell in ws[header_row]]
    for idx, h in enumerate(header_values):
        if h == "ROOM": mapping[0] = idx
        elif h == "BOQ REF": mapping[1] = idx
        elif h.endswith("CODE") and 2 not in mapping: mapping[2] = idx
        else:
            for expected_idx, expected in enumerate(COLUMN_HEADERS):
                if expected_idx > 2 and _norm(expected) == h: mapping[expected_idx] = idx
    if not all(k in mapping for k in range(4)):
        fail(422, "HEADERS_NOT_FOUND", "Unable to map compulsory Drawing Log fields")
    legacy_code_header = _cell(ws.cell(header_row, mapping[2]+1).value) or "INTERNAL ITEM CODE"
    sync = {}
    if "_OCERP_SYNC" in book.sheetnames:
        sh = book["_OCERP_SYNC"]
        for cells in sh.iter_rows(min_row=2, values_only=True):
            if cells and cells[0] is not None:
                try:
                    if len(cells) >= 10 and cells[9] is not None and int(cells[9]) != SYNC_SCHEMA:
                        fail(422, "INVALID_SYNC_SCHEMA", "Unsupported hidden Drawing Log sync schema")
                    sync[int(cells[0])] = {"project_id": cells[1], "register_id": cells[2], "item_id": cells[3],
                      "version": int(cells[4]), "row_hash": cells[5], "checkpoint_data": json.loads(cells[6] or "{}"),
                      "visible_hashes": json.loads(cells[7] or "{}"),
                      "original_identity": json.loads(cells[8] or "{}") if len(cells)>8 else {}}
                except (ValueError, TypeError, json.JSONDecodeError, IndexError):
                    fail(422, "INVALID_SYNC", "Malformed _OCERP_SYNC row")
    parsed = []
    warnings = []
    floor = category = ""
    occurrence: Counter = Counter()
    for idx in range(header_row + 1, ws.max_row + 1):
        source = list(ws[idx])
        if any(c.data_type == "f" for c in source):
            fail(422, "FORMULA_FORBIDDEN", f"Formula in row {idx}")
        values = [_cell(source[mapping[k]].value) if k in mapping and mapping[k] < len(source) else None for k in range(20)]
        if not any(x is not None and x.strip() for x in values):
            continue
        context = (values[0] or "").strip()
        no_body = not any(values[k] for k in [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 17, 18, 19])
        if no_body and context:
            if re.search(r"\b(FLOOR|PIANO|LEVEL|LIVELLO)\b", context.upper()):
                floor, category = context, ""
            else:
                category = context
            continue
        if not values[2] and not values[3]:
            warnings.append({"row": idx, "code": "MISSING_ITEM", "message": "No internal code or description"})
            continue
        if len(parsed) >= MAX_ROWS:
            fail(422, "TOO_MANY_ROWS", "Max 10,000 items")
        # Legacy workbook has one single ROOM column; inherit floor + category from context,
        # and keep actual room text per line. Never manufacture physical instances.
        room = context
        code = values[2] or ""
        k = (floor, room, code)
        occurrence[k] += 1
        try:
            original_approval = source[mapping[10]].value if 10 in mapping else None
            if isinstance(original_approval, (int, float)):
                converted = from_excel(original_approval, epoch=book.epoch)
                approval = converted.date() if isinstance(converted, datetime) else converted
                if not isinstance(approval, date):
                    raise ValueError("Non-date serial")
            else:
                approval = date.fromisoformat(values[10][:10]) if values[10] else None
        except (ValueError, OverflowError):
            approval = None
            if values[10]: warnings.append({"row": idx, "code": "INVALID_APPROVAL_DATE"})
        row = {"excel_row": idx, "floor": floor, "room": room, "category": category,
               "item_code": code, "description": values[3] or "", "specification": values[4] or "",
               "quoted_qty_raw": values[5], "production_qty_raw": values[6],
               "comments": values[7] or "", "change_notes": values[8] or "", "drafter_name": values[9],
               "reported_approval_date": approval.isoformat() if approval else None,
               "notes": values[17] or "", "boq_refs": _boq_refs(values[1]), "checkpoint_raw":
                   {ws_code: values[col] for col, ws_code in LEGACY_WORKSTREAMS.items()},
               "occurrence_index": occurrence[k], "sync": sync.get(idx)}
        row["checkpoint_structured"] = {}
        if row["sync"]:
            # The value shown in the Excel cell may be the status, not the raw
            # text. An unchanged cell preserves the structured record verbatim.
            saved = row["sync"]["checkpoint_data"]
            hashes = row["sync"]["visible_hashes"]
            for code, raw in list(row["checkpoint_raw"].items()):
                if code in saved and hashes.get(code) == _hash(raw):
                    if not isinstance(saved[code], dict):
                        fail(422, "INVALID_SYNC", "Invalid checkpoint payload")
                    row["checkpoint_raw"][code] = saved[code].get("reported_value")
                    row["checkpoint_structured"][code] = saved[code]
        parsed.append(row)
    return {"rows": parsed, "warnings": warnings, "sha256": hashlib.sha256(content).hexdigest(),
            "legacy_code_header": legacy_code_header, "source_sheet": ws.title}


def needs_checkpoint_update(existing: ItemCheckpoint | None, raw: str | None) -> bool:
    """A cleared visible cell must clear an existing raw value, not be ignored."""
    return (existing is not None and existing.reported_value != raw) or (existing is None and raw is not None)


def restored_checkpoint_fields(row: dict, code: str) -> dict:
    """The hidden worksheet is only a lossless snapshot, never an owner of state.

    For a new entry from a same-register workbook, restore explicitly recorded
    checkpoint fields; for updates the current database record wins. All values
    go through Pydantic, so a malformed hidden sheet cannot store arbitrary status.
    """
    saved = row.get("checkpoint_structured", {}).get(code)
    if not saved:
        return {}
    allowed = {key: saved[key] for key in CheckpointWrite.model_fields
               if key in saved and key != "expected_version"}
    try:
        return CheckpointWrite(expected_version=0, **allowed).model_dump(exclude={"expected_version"})
    except Exception:
        fail(422, "INVALID_SYNC", "Malformed structured checkpoint in hidden sync worksheet")


def reference_warnings(rows: list[dict], positions: list[dict]) -> list[dict]:
    """Explain unresolved and duplicated commercial ordinals without guessing.

    A reference may validly occur on many ProductionItems; the ambiguity is
    specifically more than one position with the same ordinal in the baseline.
    """
    counts = Counter(str(position.get("ordinal", "")) for position in positions)
    warnings: list[dict] = []
    for row in rows:
        for reference in row["boq_refs"]:
            count = counts.get(reference, 0)
            if count != 1:
                warnings.append({
                    "row": row["excel_row"],
                    "reference": reference,
                    "code": "AMBIGUOUS_BOQ_REF" if count > 1 else "UNRESOLVED_BOQ_REF",
                    "message": "BOQ reference needs explicit mapping" if count > 1 else "Not present in the approved BOQ baseline",
                })
    return warnings


def _item_values(item: ProductionItem) -> dict:
    fields = ("floor", "room", "category", "item_code", "description", "specification",
              "quoted_qty_raw", "production_qty_raw", "comments", "change_notes", "drafter_name", "notes")
    return {f: getattr(item, f) for f in fields} | {"reported_approval_date": _cell(item.reported_approval_date)}

async def build_plan(s: RegisterService, reg, parsed: dict, resolutions: dict | None = None) -> dict:
    resolutions = resolutions or {}
    current = list((await s.db.execute(select(ProductionItem).where(ProductionItem.register_id == reg.id,
                  ProductionItem.lifecycle != "archived").order_by(ProductionItem.display_order, ProductionItem.created_at))).scalars())
    by_id = {str(x.id): x for x in current}
    keyed: dict[tuple, list] = defaultdict(list)
    for x in current:
        keyed[x.floor, x.room, x.item_code].append(x)
    result = []
    for row in parsed["rows"]:
        sync = row.get("sync")
        matching = None
        issue = None
        if sync:
            if sync["project_id"] != str(s.project_id) or sync["register_id"] != str(reg.id):
                issue = "WRONG_REGISTER"
            elif sync["item_id"] not in by_id:
                issue = "UNKNOWN_ITEM_UUID"
            else:
                matching = by_id[sync["item_id"]]
                if matching.version != sync["version"]:
                    issue = "ITEM_VERSION_CONFLICT"
                identity = sync.get("original_identity") or {}
                if identity and any(str(row[key] or "") != str(identity.get(key) or "") for key in ("floor", "room", "item_code")):
                    issue = "SYNC_IDENTITY_CHANGED"
        else:
            candidates = keyed.get((row["floor"], row["room"], row["item_code"]), [])
            if len(candidates) == 1 and row["occurrence_index"] == 1:
                matching = candidates[0]
            elif candidates:
                issue = "AMBIGUOUS_LEGACY_MATCH"
        override = resolutions.get(str(row["excel_row"]))
        if override is not None:
            validate_resolution_override(issue, override)
            if not isinstance(override, dict) or override.get("action") not in ("create", "match"):
                fail(422, "INVALID_RESOLUTION", "Expected create or match action")
            if override["action"] == "create":
                matching, issue = None, None
            else:
                selected = override.get("item_id")
                if selected not in by_id:
                    fail(422, "INVALID_RESOLUTION", "Selected item is not in this register")
                matching, issue = by_id[selected], None
        vals = {k: row[k] for k in _item_values(matching) if k in row} if matching else {}
        changed = any(_cell(getattr(matching, k)) != _cell(v) for k, v in vals.items()) if matching else True
        current_links = [x.boq_ref_raw for x in await s.repo.links(matching.id)] if matching else []
        changed = changed or row["boq_refs"] != current_links
        if matching:
            recorded = {x.workstream_id: x for x in await s.repo.checkpoints(matching.id)}
            for workstream in await s.repo.workstreams(reg.id):
                if workstream.code in row["checkpoint_raw"]:
                    present = recorded.get(workstream.id)
                    if (_cell(present.reported_value) if present else None) != row["checkpoint_raw"][workstream.code]:
                        changed = True
        if issue == "ITEM_VERSION_CONFLICT" and not changed:
            issue = None  # Safe idempotent retry: file describes the current value.
        if not matching and not issue:
            action = "create"
        elif issue:
            action = "conflict"
        elif changed:
            action = "update"
        else:
            action = "unchanged"
        result.append({"row": row["excel_row"], "item_id": str(matching.id) if matching else None,
                       "item_version": matching.version if matching else None, "action": action,
                       "issue": issue, "values": {k: _cell(v) for k, v in row.items() if k not in ("sync", "checkpoint_raw", "checkpoint_structured")},
                       "boq_refs": row["boq_refs"]})
    return {"items": result, "counts": dict(Counter(x["action"] for x in result))}

async def preview(s: RegisterService, reg, content: bytes, filename: str) -> dict:
    parsed = parse_xlsx(content, filename)
    plan = await build_plan(s, reg, parsed)
    warnings = [*parsed["warnings"], *reference_warnings(parsed["rows"], await s.boq_positions(reg))]
    return {"preview_sha256": parsed["sha256"], "register_version": reg.version,
            "legacy_code_header": parsed["legacy_code_header"], "warnings": warnings, **plan}

async def commit(s: RegisterService, reg, content: bytes, filename: str, preview_sha256: str,
                 expected_register_version: int, resolutions_json: str) -> dict:
    parsed = parse_xlsx(content, filename)
    if parsed["sha256"] != preview_sha256:
        fail(409, "FILE_CHANGED", "The uploaded file differs from its preview")
    try:
        resolutions = json.loads(resolutions_json)
    except json.JSONDecodeError:
        fail(422, "INVALID_RESOLUTIONS", "Invalid resolutions JSON")
    if not isinstance(resolutions, dict):
        fail(422, "INVALID_RESOLUTIONS", "Resolutions must be an object")
    plan = await build_plan(s, reg, parsed, resolutions=resolutions)
    used = [p["item_id"] for p in plan["items"] if p["item_id"]]
    if len(set(used)) != len(used):
        fail(422, "DUPLICATE_ITEM_MATCH", "Several spreadsheet rows target the same OCERP item")
    if reg.version != expected_register_version and any(x["action"] != "unchanged" for x in plan["items"]):
        fail(409, "VERSION_CONFLICT", "Register changed after preview")
    ws_by_code = {w.code: w for w in await s.repo.workstreams(reg.id)}
    changed = 0
    # Validate the complete plan before the first mutation, including legacy ambiguity.
    for entry in plan["items"]:
        if entry["action"] == "conflict":
            fail(409, "IMPORT_CONFLICT", f"Row {entry['row']}: {entry['issue']}; explicit resolution required")
    for row, entry in zip(parsed["rows"], plan["items"], strict=True):
        if entry["action"] == "unchanged":
            continue
        vals = {k: v for k, v in row.items() if k in ItemInput.model_fields and k != "boq_links"}
        vals["display_order"] = row["excel_row"]
        vals["boq_links"] = [BOQLinkInput(boq_ref_raw=x) for x in row["boq_refs"]]
        if entry["action"] == "create":
            item = await s.add_item(reg, ItemInput(**vals), origin="excel")
            changed += 1
        else:
            item = await s.require_item(reg, UUID(entry["item_id"]))
            s.check_version(item.version, entry["item_version"])
            before = serialize(item)
            for key in vals:
                if key != "boq_links" and key != "display_order":
                    value = vals[key]
                    if key == "reported_approval_date" and isinstance(value, str):
                        value = date.fromisoformat(value)
                    setattr(item, key, value)
            await s.replace_links(reg, item, vals["boq_links"])
            item.version += 1
            await s.bump(reg)
            s.audit(reg, "item", item.id, "import_updated", before=before, after=serialize(item), source="excel")
            changed += 1
        for code, raw in row["checkpoint_raw"].items():
            ws = ws_by_code.get(code)
            if ws is None:
                continue
            existing = (await s.db.execute(select(ItemCheckpoint).where(ItemCheckpoint.item_id == item.id,
                        ItemCheckpoint.workstream_id == ws.id))).scalar_one_or_none()
            restored = restored_checkpoint_fields(row, code) if existing is None else {}
            if not needs_checkpoint_update(existing, raw) and not restored:
                continue
            if existing is None:
                existing = ItemCheckpoint(item_id=item.id, project_id=s.project_id, workstream_id=ws.id,
                                         status="unknown", version=1, source="excel")
                for field, value in restored.items():
                    setattr(existing, field, value)
                s.db.add(existing)
            else:
                existing.version += 1
            existing.reported_value = raw
            existing.source = "excel"
            existing.updated_by = s.actor_id
            await s.db.flush()
            s.audit(reg, "checkpoint", existing.id, "import_updated", after=serialize(existing), source="excel")
    if changed:
        if reg.legacy_code_header == "INTERNAL ITEM CODE":
            reg.legacy_code_header = parsed["legacy_code_header"]
        await s.db.flush()
        s.event(reg, "import.completed")
    return {"created": plan["counts"].get("create", 0), "updated": plan["counts"].get("update", 0),
            "unchanged": plan["counts"].get("unchanged", 0), "register_version": reg.version}

async def export_xlsx(s: RegisterService, reg, mode: str) -> bytes:
    items = list((await s.db.execute(select(ProductionItem).where(ProductionItem.register_id == reg.id,
              ProductionItem.lifecycle != "archived").order_by(ProductionItem.display_order, ProductionItem.created_at))).scalars())
    book = Workbook()
    ws = book.active
    ws.title = "Foglio1"
    ws.cell(1, 1, f"{s.project_id} - DRAWING LOG")
    ws.append([])
    for col, name in enumerate(COLUMN_HEADERS, start=1):
        _write_text(ws.cell(3, col), reg.legacy_code_header if col == 3 else name)
        ws.cell(3, col).font = Font(bold=True, color="FFFFFF")
        ws.cell(3, col).fill = PatternFill("solid", fgColor="273E58")
    meta = book.create_sheet("_OCERP_SYNC")
    meta.append(["row_number", "project_id", "register_id", "item_id", "item_version", "row_hash", "checkpoint_data", "visible_hashes", "original_identity", "schema_version"])
    meta.sheet_state = "hidden"
    workstreams = {w.id: w for w in await s.repo.workstreams(reg.id)}
    prev_floor = prev_category = None
    for item in items:
        if item.floor != prev_floor:
            if item.floor:
                ws.append([item.floor])
            prev_floor, prev_category = item.floor, None
        if item.category != prev_category:
            if item.category:
                ws.append([item.category])
            prev_category = item.category
        links = [l.boq_ref_raw for l in await s.repo.links(item.id)]
        checkpoints = await s.repo.checkpoints(item.id)
        by_code = {workstreams[c.workstream_id].code: c for c in checkpoints if c.workstream_id in workstreams}
        visible = {code: c.reported_value if c.reported_value is not None else (c.status if c.status != "unknown" else None)
                   for code, c in by_code.items()}
        values = [item.room, "\n".join(links), item.item_code, item.description, item.specification,
                  item.quoted_qty_raw, item.production_qty_raw, item.comments, item.change_notes, item.drafter_name,
                  _cell(item.reported_approval_date)] + [None] * 9
        values[17] = item.notes
        for col, code in LEGACY_WORKSTREAMS.items():
            values[col] = visible.get(code)
        row_index = ws.max_row + 1
        for col, value in enumerate(values, 1):
            _write_text(ws.cell(row_index, col), value)
        meta.append([row_index, str(s.project_id), str(reg.id), str(item.id), item.version,
                     _hash(values), json.dumps({k: serialize(c) for k, c in by_code.items()}, default=str),
                     json.dumps({k: _hash(v) for k, v in visible.items()}),
                     json.dumps({"floor": item.floor, "room": item.room, "item_code": item.item_code}), SYNC_SCHEMA])
    ws.freeze_panes = "D4"
    ws.auto_filter.ref = f"A3:T{ws.max_row}"
    for col, width in zip("ABCDEFGHIJKLMNOPQRST", [24, 29, 24, 45, 35, 17, 19, 35, 30, 23, 20, 20, 20, 20, 22, 18, 22, 28, 23, 23]):
        ws.column_dimensions[col].width = width
    if mode == "extended":
        all_items = book.create_sheet("Items")
        all_items.append(list(serialize(items[0]).keys()) if items else ["id"])
        for item in items: all_items.append([str(v) if v is not None else "" for v in serialize(item).values()])
        for sheet_name, content in (("BOQ links", [(l.item_id, l.boq_ref_raw, l.boq_position_id, l.resolution) for item in items for l in await s.repo.links(item.id)]),
                                    ("Workstreams", [(w.id, w.code, w.label, w.enabled) for w in workstreams.values()]),
                                    ("Checkpoints", [(c.item_id, workstreams[c.workstream_id].code, c.status, c.reported_value,
                                       c.planned_date, c.actual_date, c.assignee_name, c.notes, c.version, c.source)
                                       for item in items for c in await s.repo.checkpoints(item.id) if c.workstream_id in workstreams]),
                                    ("References", [(r.item_id, r.kind, r.owner_id, r.namespace, r.external_key, r.uri, r.label)
                                       for item in items for r in await s.repo.references(item.id)])):
            sheet = book.create_sheet(sheet_name)
            for row in content: sheet.append([str(x) if x is not None else "" for x in row])
    output = BytesIO()
    book.save(output)
    return output.getvalue()
